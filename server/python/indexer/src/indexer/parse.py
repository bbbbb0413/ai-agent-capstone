"""형식별 파서: 원본 파일에서 본문과 제목 구조를 뽑는다 (SPEC-document-store 5.3).

ai-web-service index-worker의 parse.py를 옮겼다(contracts/indexing/README.md §3). 모든 파서가
"본문 + 제목 목록(수준, 제목, 위치)"을 돌려주고, 공통 청커(chunk.py)가 제목 경로로 절을 나눈다.
서식은 재현하지 않는다. 검색과 인용에 필요한 것은 글자와 구조다.
OpsLens에서 바꾼 점:
- Markdown 코드 펜스 안의 `#` 줄은 제목으로 보지 않는다 (원본의 알려진 한계)
- PDF는 글자 크기 추정에 더해 `제N조`·`RB-STD-0N`으로 시작하는 줄도 제목으로 본다 (계약서·표준 런북의 조항)
- Word 제목은 스타일 이름(`Heading 1`, 한국어 Word의 `제목 1`)으로 찾는다
- pptx·xlsx는 범위 밖이다 (SPEC 11장)
"""

import io
import re
from dataclasses import dataclass

import docx
import pdfplumber
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from pdfminer.pdfdocument import PDFEncryptionError, PDFPasswordIncorrect

FORMATS = ("md", "txt", "pdf", "docx")

# 실제 계약서·런북이 이보다 길 일은 없다. 넘으면 거절한다
MAX_PDF_PAGES = 200
_MAX_CELL_CHARS = 500

_MD_HEADING = re.compile(r"^(#{1,6})[ \t]+(\S.*?)[ \t#]*$")
_MD_FENCE = re.compile(r"^\s{0,3}(```|~~~)")
_HEADING_STYLE = re.compile(r"(?:heading|제목)\s*([1-6])", re.IGNORECASE)
# 계약서 조항(제5조, 제 5 조)과 표준 런북 조항(RB-STD-02). 글자 크기가 본문과 같아도 제목으로 본다
_NUMBERED_TITLE = re.compile(r"^(?:제\s*\d+\s*조|RB-STD-\d{2})(?![\w-])")


@dataclass(frozen=True)
class Heading:
    level: int  # 1이 최상위
    title: str
    offset: int  # text 안에서 이 제목 줄이 시작하는 위치


@dataclass(frozen=True)
class ParsedDocument:
    text: str
    headings: tuple[Heading, ...] = ()


class ParseError(Exception):
    """파싱 실패. retryable이 False면 다시 시도해도 결과가 같다(암호화·손상·미지원 형식)."""

    def __init__(self, message: str, *, retryable: bool = False):
        super().__init__(message)
        self.retryable = retryable


def parse_document(data: bytes, file_format: str) -> ParsedDocument:
    if file_format == "md":
        return parse_markdown(_utf8(data))
    if file_format == "txt":
        # 일반 텍스트는 제목 표기가 없다. 짧은 줄을 소제목으로 짐작하면 분할이 어긋나므로 길이로만 나눈다
        return ParsedDocument(_utf8(data))
    if file_format == "pdf":
        return parse_pdf(data)
    if file_format == "docx":
        return parse_docx(data)
    raise ParseError(f"지원하지 않는 파일 형식이다: {file_format}")


def _utf8(data: bytes) -> str:
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        raise ParseError("UTF-8 텍스트가 아니다") from None


def parse_markdown(text: str) -> ParsedDocument:
    headings: list[Heading] = []
    offset, fence = 0, None
    for line in text.splitlines(keepends=True):
        stripped = line.rstrip("\r\n")
        opener = _MD_FENCE.match(stripped)
        if opener:
            fence = None if fence == opener.group(1) else (fence or opener.group(1))
        elif fence is None and (m := _MD_HEADING.match(stripped)):
            headings.append(Heading(level=len(m.group(1)), title=m.group(2).strip(), offset=offset))
        offset += len(line)
    return ParsedDocument(text, tuple(headings))


# ---- 표 ----
# 표는 마크다운 파이프 행(`| 셀 | 셀 |`)으로 적는다. 청커가 이 모양의 줄은 길이 분할에서 끊지 않는다.
# 머리글 구분선(`|---|`)은 붙이지 않는다. 첫 행이 머리글인지 파일에서 알 수 없다.

def _cell_text(value: object) -> str:
    text = "" if value is None else str(value)
    text = text.replace("\r", " ").replace("\n", " ").replace("|", r"\|").strip()
    return text[:_MAX_CELL_CHARS]


def _table_row(cells) -> str:
    rendered = [_cell_text(cell) for cell in cells]
    while rendered and not rendered[-1]:
        rendered.pop()
    return "| " + " | ".join(rendered) + " |" if rendered else ""


# ---- PDF ----
# PDF에는 "이건 제목이다"라는 표시가 없어 글자 크기로 추정한다. 확신이 없으면 제목을 만들지 않는다.
# 제목을 잘못 만들면 인용이 문서의 엉뚱한 곳을 가리키는데, 그건 구조가 없는 것보다 나쁘다.
_PDF_HEADING_SIZE_RATIO = 1.15   # 본문 크기의 이 배 이상
_PDF_HEADING_MAX_CHARS = 120     # 제목은 짧다
_PDF_BODY_MIN_SHARE = 0.5        # 한 크기가 글자 수의 과반이어야 그것을 본문으로 믿는다
_MAX_LEVEL = 6


@dataclass(frozen=True)
class PdfLine:
    text: str
    offset: int
    size: float  # 그 줄 글자 크기의 중앙값


def _median_size(chars: list) -> float:
    sizes = sorted(float(c["size"]) for c in chars if c.get("size") is not None)
    if not sizes:
        return 0.0
    middle = len(sizes) // 2
    return sizes[middle] if len(sizes) % 2 else (sizes[middle - 1] + sizes[middle]) / 2


def infer_pdf_headings(lines: list[PdfLine], weights: dict[float, int]) -> tuple[Heading, ...]:
    """글자 크기로 추정한 제목에 번호 패턴(제N조, RB-STD-0N) 줄을 더한다."""
    found: dict[int, Heading] = {}
    total = sum(weights.values())
    body = max(weights.items(), key=lambda kv: (kv[1], kv[0]))[0] if weights else 0.0
    if body > 0 and len(lines) >= 3 and total > 0 and weights.get(body, 0) / total >= _PDF_BODY_MIN_SHARE:
        candidates = [line for line in lines
                      if line.size >= body * _PDF_HEADING_SIZE_RATIO
                      and 0 < len(line.text.strip()) <= _PDF_HEADING_MAX_CHARS]
        ranks = sorted({line.size for line in candidates}, reverse=True)
        level_of = {size: min(i + 1, _MAX_LEVEL) for i, size in enumerate(ranks)}
        for line in candidates:
            found[line.offset] = Heading(level_of[line.size], line.text.strip(), line.offset)
    numbered_level = min(max((h.level for h in found.values()), default=0) + 1, _MAX_LEVEL)
    for line in lines:
        title = line.text.strip()
        if line.offset not in found and len(title) <= _PDF_HEADING_MAX_CHARS and _NUMBERED_TITLE.match(title):
            found[line.offset] = Heading(numbered_level, title, line.offset)
    return tuple(found[offset] for offset in sorted(found))


def parse_pdf(data: bytes) -> ParsedDocument:
    try:
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            if len(pdf.pages) > MAX_PDF_PAGES:
                raise ParseError(f"PDF 쪽 수가 너무 많다({len(pdf.pages)}쪽, 최대 {MAX_PDF_PAGES}쪽)")
            parts: list[str] = []
            lines: list[PdfLine] = []
            weights: dict[float, int] = {}
            offset = 0
            for page in pdf.pages:
                for line in page.extract_text_lines(layout=False) or []:
                    text = line["text"]
                    chars = line.get("chars") or []
                    size = _median_size(chars)
                    lines.append(PdfLine(text, offset, size))
                    weights[size] = weights.get(size, 0) + len(chars)
                    parts.append(text + "\n")
                    offset += len(text) + 1
                parts.append("\n")   # 쪽 경계
                offset += 1
            text = "".join(parts)
    except ParseError:
        raise
    except Exception as e:  # noqa: BLE001 — 손상 파일은 라이브러리마다 다른 예외를 던진다
        if _is_encrypted(e):
            raise ParseError("암호화된 PDF라 본문을 추출할 수 없다") from None
        raise ParseError(f"PDF를 읽을 수 없다: {type(e).__name__}") from None
    if not text.strip():
        raise ParseError("텍스트가 없는 PDF다 (스캔 문서는 지원하지 않는다)")
    return ParsedDocument(text, infer_pdf_headings(lines, weights))


def _is_encrypted(error: BaseException) -> bool:
    """pdfplumber는 pdfminer 예외를 감싸서 던진다. 원인 사슬을 따라가며 확인한다."""
    seen: set[int] = set()
    pending: list[object] = [error]
    while pending:
        current = pending.pop()
        if not isinstance(current, BaseException) or id(current) in seen:
            continue
        seen.add(id(current))
        if isinstance(current, (PDFPasswordIncorrect, PDFEncryptionError)):
            return True
        pending.extend([current.__cause__, current.__context__, *current.args])
    return False


# ---- Word ----

def _docx_blocks(document):
    """문단과 표를 문서에 놓인 순서대로 준다. 표가 어느 제목 아래인지가 절 경로를 정한다."""
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, document)
        elif child.tag == qn("w:tbl"):
            yield Table(child, document)


def _docx_row_cells(row) -> list[str]:
    """한 행의 셀 글자. 가로로 병합된 칸은 한 번만 센다."""
    values, previous = [], None
    for cell in row.cells:
        if cell._tc is previous:
            continue
        previous = cell._tc
        values.append(cell.text)
    return values


def _heading_level(paragraph: Paragraph) -> int | None:
    try:
        name = paragraph.style.name if paragraph.style is not None else ""
    except (KeyError, ValueError):
        name = ""
    match = _HEADING_STYLE.search(name or "")
    return int(match.group(1)) if match else None


def parse_docx(data: bytes) -> ParsedDocument:
    """본문 문단과 표를 읽는다. 머리글·바닥글·각주·텍스트 상자는 빠진다."""
    try:
        document = docx.Document(io.BytesIO(data))
        parts: list[str] = []
        headings: list[Heading] = []
        length = 0
        for block in _docx_blocks(document):
            if isinstance(block, Table):
                for row in block.rows:
                    line = _table_row(_docx_row_cells(row))
                    if line:
                        parts.append(line + "\n")
                        length += len(line) + 1
                continue
            line = block.text or ""
            if not line.strip():
                parts.append("\n")
                length += 1
                continue
            level = _heading_level(block)
            if level is not None:
                headings.append(Heading(level, line.strip(), length))
            parts.append(line + "\n")
            length += len(line) + 1
    except Exception as e:  # noqa: BLE001
        raise ParseError(f"docx를 읽을 수 없다: {type(e).__name__}") from None
    text = "".join(parts)
    if not text.strip():
        raise ParseError("본문이 비어 있는 docx다")
    return ParsedDocument(text, tuple(headings))
