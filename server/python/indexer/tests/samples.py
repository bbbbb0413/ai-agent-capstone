"""테스트용 원본 파일을 만든다. 저장소에 바이너리 샘플을 두지 않는다."""

import io

import docx


def pdf(lines: list[tuple[float, str]], pages: int = 1) -> bytes:
    """Helvetica로 (글자 크기, 문장) 줄을 쓴 최소 PDF. 문장은 ASCII만 된다(글꼴을 넣지 않는다)."""
    def escape(text: str) -> str:
        return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")

    contents = []
    for _ in range(pages):
        y, ops = 800, ["BT"]
        for size, text in lines:
            ops.append(f"/F1 {size} Tf 1 0 0 1 50 {y} Tm ({escape(text)}) Tj")
            y -= int(size * 1.8)
        ops.append("ET")
        contents.append("\n".join(ops).encode())

    objects: list[bytes] = []
    page_ids = [3 + 2 * i for i in range(pages)]
    font_id = 3 + 2 * pages
    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objects.append(f"<< /Type /Pages /Kids [{' '.join(f'{p} 0 R' for p in page_ids)}] /Count {pages} >>".encode())
    for i, content in enumerate(contents):
        objects.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 842] /Contents {page_ids[i] + 1} 0 R "
                       f"/Resources << /Font << /F1 {font_id} 0 R >> >> >>".encode())
        objects.append(b"<< /Length %d >>\nstream\n" % len(content) + content + b"\nendstream")
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")

    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{number} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets:
        out.write(f"{offset:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return out.getvalue()


def blank_pdf() -> bytes:
    """글자가 없는 PDF (스캔 문서처럼 텍스트 층이 없다)."""
    return pdf([])


def docx_file(blocks: list[tuple[str, object]]) -> bytes:
    """("h1"·"h2"·"p", 글) 또는 ("table", [[셀, …], …]) 목록으로 Word 문서를 만든다."""
    document = docx.Document()
    for kind, value in blocks:
        if kind.startswith("h"):
            document.add_heading(str(value), level=int(kind[1:]))
        elif kind == "p":
            document.add_paragraph(str(value))
        elif kind == "table":
            rows = list(value)
            table = document.add_table(rows=len(rows), cols=len(rows[0]))
            for r, row in enumerate(rows):
                for c, cell in enumerate(row):
                    table.cell(r, c).text = str(cell)
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()
